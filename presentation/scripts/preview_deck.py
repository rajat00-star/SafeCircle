"""Approximate slide previews and a text-overflow check for the deck, for machines without PowerPoint,
Keynote or LibreOffice. It is NOT a faithful renderer: shapes, pictures and text are drawn with Pillow,
and Calibri is measured with Arial (wider), so the overflow check errs on the side of reporting too much.

    python preview_deck.py presentation/SafeCircle-Guardian-PS02.pptx <out-dir>
Requires python-pptx and Pillow (use a throwaway virtual environment).
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

FONT_DIR = Path("/System/Library/Fonts/Supplemental")
DPI = 96
SCHEME = {"BACKGROUND_2": "E9F0FB", "ACCENT_1": "3EE6FF", "ACCENT_2": "A08CFF", "ACCENT_3": "FF5468", "ACCENT_4": "5CEBB0",
          "ACCENT_5": "FFB84D", "ACCENT_6": "9AA8C2", "TEXT_1": "05070D", "BACKGROUND_1": "FFFFFF"}


def px(emu) -> int:
    return int(Emu(emu).inches * DPI)


def font(size_pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_DIR / ("Arial Bold.ttf" if bold else "Arial.ttf")), max(6, int(size_pt * DPI / 72)))


def color_of(run_font, default="E9F0FB") -> str:
    try:
        c = run_font.color
        if c and c.type is not None:
            if c.rgb is not None:
                return str(c.rgb)
            if c.theme_color is not None:
                return SCHEME.get(c.theme_color.name, default)
    except Exception:
        pass
    return default


def wrap_height(draw, paragraphs, width_px, default_size) -> int:
    total = 0
    for p in paragraphs:
        runs = [r for r in p.runs if r.text]
        size = next((r.font.size.pt for r in runs if r.font.size), default_size)
        bold = any(r.font.bold for r in runs)
        f = font(size, bold)
        words = "".join(r.text for r in runs).split(" ")
        line, lines = "", 1
        for w in words:
            trial = (line + " " + w).strip()
            if draw.textlength(trial, font=f) > width_px and line:
                lines += 1
                line = w
            else:
                line = trial
        total += int(lines * size * 1.2 * DPI / 72) + 3
    return total


def main(path: str, out: str) -> None:
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    Path(out).mkdir(parents=True, exist_ok=True)
    problems = []
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), "#05070D")
        d = ImageDraw.Draw(img)
        bg = slide.background.fill
        try:
            if bg.type == 6:  # picture background
                pass
        except Exception:
            pass
        # Picture backgrounds are stored as a blip on the slide's <p:bg>; draw them first if present.
        for rel in slide.part.rels.values():
            if "image" in rel.reltype and 'r:embed="%s"' % rel.rId in slide._element.xml.split("<p:spTree>")[0]:
                img.paste(Image.open(io.BytesIO(rel.target_part.blob)).convert("RGB").resize((W, H)))
        shapes = list(slide.placeholders) + [s for s in slide.shapes if not s.is_placeholder]
        for s in slide.shapes:
            x, y, w, h = px(s.left or 0), px(s.top or 0), px(s.width or 0), px(s.height or 0)
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img.paste(Image.open(io.BytesIO(s.image.blob)).convert("RGB").resize((max(1, w), max(1, h))), (x, y))
                continue
            if s.shape_type == MSO_SHAPE_TYPE.CHART:
                d.rectangle([x, y, x + w, y + h], outline="#3EE6FF")
                d.text((x + 6, y + 6), "[native chart]", fill="#9AA8C2", font=font(10, False))
                continue
            if s.shape_type == MSO_SHAPE_TYPE.TABLE:
                tbl = s.table
                cy = y
                for r in tbl.rows:
                    rh = px(r.height)
                    cx = x
                    for ci, cell in enumerate(r.cells):
                        cw = px(tbl.columns[ci].width)
                        d.rectangle([cx, cy, cx + cw, cy + rh], outline="#22324F", fill="#0F1830")
                        need = wrap_height(d, cell.text_frame.paragraphs, cw - 10, 12)
                        if need > rh + 4:
                            problems.append(f"slide {i}: table cell '{cell.text[:40]}' needs {need}px, row is {rh}px (row grows in PowerPoint)")
                        d.multiline_text((cx + 5, cy + 3), cell.text[:60], fill="#E9F0FB", font=font(9, False))
                        cx += cw
                    cy += rh
                continue
            try:
                fill = s.fill
                if fill.type == 1:
                    rgb = str(fill.fore_color.rgb) if fill.fore_color.type == 1 else "0F1830"
                    d.rounded_rectangle([x, y, x + w, y + h], radius=6, fill="#" + rgb)
                if s.line and s.line.fill.type == 1:
                    d.rounded_rectangle([x, y, x + w, y + h], radius=6, outline="#22324F")
            except Exception:
                pass
            if s.has_text_frame and s.text_frame.text.strip():
                tf = s.text_frame
                need = wrap_height(d, tf.paragraphs, w, 14)
                if need > h + 6:
                    problems.append(f"slide {i}: '{tf.text[:50]}' needs ~{need}px, box is {h}px")
                ty = y
                for p in tf.paragraphs:
                    runs = [r for r in p.runs if r.text]
                    if not runs:
                        continue
                    size = next((r.font.size.pt for r in runs if r.font.size), 14)
                    f = font(size, any(r.font.bold for r in runs))
                    txt, line = "".join(r.text for r in runs), ""
                    col = "#" + color_of(runs[0].font)
                    for word in txt.split(" "):
                        trial = (line + " " + word).strip()
                        if d.textlength(trial, font=f) > w and line:
                            d.text((x, ty), line, fill=col, font=f)
                            ty += int(size * 1.2 * DPI / 72)
                            line = word
                        else:
                            line = trial
                    d.text((x, ty), line, fill=col, font=f)
                    ty += int(size * 1.2 * DPI / 72) + 3
                if x + w > W + 2 or y + h > H + 2:
                    problems.append(f"slide {i}: '{tf.text[:40]}' extends past the slide edge")
        img.save(Path(out) / f"slide-{i:02d}.png")
    print("\n".join(problems) or "No overflow found by the approximate check.")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
