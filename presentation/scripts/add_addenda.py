"""Append a dated addendum to the research paper and the idea document (2026-10-12).

The original documents (2026-10-10) predate the PS02 case engine and its synthetic validation. Rather than
rewrite them, this appends one clearly dated section listing what changed, with figures read from
presentation/content/facts.json. Run once: python add_addenda.py (needs python-docx).
"""
import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
F = json.loads((ROOT / "presentation/content/facts.json").read_text())
v2, v3 = F["v2"], F["v3_after"]
s = v2["safecircle"]
MARK = "Addendum (2026-10-12)"


def add(path: Path, short: bool) -> None:
    doc = Document(str(path))
    if any(p.text.startswith(MARK) for p in doc.paragraphs):
        print(f"{path.name}: addendum already present"); return
    doc.add_page_break()
    doc.add_heading(f"{MARK}: changes since this document was written", level=1)
    paras = [
        "This document was written on 2026-10-10. Since then the product gained a case layer for Problem Statement 02 "
        "(AI-driven scam pattern recognition): Scam DNA investigations that re-assess a case as evidence arrives, alert "
        "de-duplication, Campaign Radar, Scam Forensics with local OCR, UPI Scam Guard, an evidence report and a scam "
        "network graph. The single-message detector described above is unchanged and its public-data results still apply to it.",
        f"The case engine was evaluated on a synthetic holdout set (v2) of {v2['scenarios']} scenarios, frozen with a checksum "
        f"before the first run. On the {v2['binary_scored']} labelled scam/benign scenarios: precision {s['precision']['value']:.3f} "
        f"({s['precision']['numerator']}/{s['precision']['denominator']}), recall {s['recall']['value']:.3f} "
        f"({s['recall']['numerator']}/{s['recall']['denominator']}), false-positive rate {s['fpr']['value']:.3f} "
        f"({s['fpr']['numerator']}/{s['fpr']['denominator']}). Errors: "
        + "; ".join(f"{e['description']} ({e['outcome']})" for e in v2["errors"]) + ".",
        f"A later v3 check ({v3['binary_scored']} scenarios, not independent) gave precision {v3['safecircle']['precision']['value']:.3f}, "
        f"recall {v3['safecircle']['recall']['value']:.3f} and false-positive rate {v3['safecircle']['fpr']['value']:.3f}. "
        "All scenarios were written by the team that wrote the rules. Synthetic benchmark performance does not establish "
        "real-world fraud-detection accuracy.",
        "Details: docs/scam-dna.md, docs/validation-report.md, docs/hardening-report.md, evaluation/holdout_v2 and "
        "evaluation/holdout_v3. Current pitch deck and 3D walkthrough: presentation/README.md.",
    ]
    for t in (paras[:1] + paras[3:] if short else paras):
        doc.add_paragraph(t)
    doc.save(str(path))
    print(f"{path.name}: addendum appended")


add(ROOT / "paper/SafeCircle-Guardian-Paper.docx", short=False)
add(ROOT / "docs/submission/SafeCircle-Guardian-Idea.docx", short=False)
